"""File reading tools — supports text, PDF, DOCX, XLSX, JSON, YAML, code files."""
import json
import os
from pathlib import Path
from typing import Optional

from jailson.config.settings import SUPPORTED_TEXT_EXTENSIONS


MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB
MAX_TEXT_CHARS = 8_000


def read_file(path: str, max_chars: int = MAX_TEXT_CHARS) -> dict:
    """Read a file and return its content as text.

    Returns: {success, content, size_bytes, extension, error}
    """
    p = Path(path).expanduser().resolve()
    if not p.exists():
        return {"success": False, "error": f"Ficheiro não encontrado: {path}"}
    if not p.is_file():
        return {"success": False, "error": f"Não é um ficheiro: {path}"}

    size = p.stat().st_size
    if size > MAX_FILE_SIZE:
        return {"success": False, "error": f"Ficheiro demasiado grande: {size / 1024 / 1024:.1f} MB (máx 5 MB)"}

    ext = p.suffix.lower()

    try:
        if ext in SUPPORTED_TEXT_EXTENSIONS:
            content = p.read_text(encoding="utf-8", errors="replace")
        elif ext == ".pdf":
            content = _read_pdf(p)
        elif ext in (".docx", ".doc"):
            content = _read_docx(p)
        elif ext in (".xlsx", ".xls"):
            content = _read_excel(p)
        elif ext in (".yaml", ".yml"):
            content = _read_yaml(p)
        else:
            try:
                content = p.read_text(encoding="utf-8", errors="replace")
            except Exception:
                return {"success": False, "error": f"Tipo de ficheiro não suportado: {ext}"}

        if len(content) > max_chars:
            content = content[:max_chars] + f"\n\n[... truncado. Total: {len(content)} caracteres]"

        return {
            "success": True,
            "content": content,
            "size_bytes": size,
            "extension": ext,
            "path": str(p),
        }
    except Exception as e:
        return {"success": False, "error": str(e)}


def list_directory(path: str, show_hidden: bool = False, recursive: bool = False, max_items: int = 200) -> dict:
    """List directory contents.

    Returns: {success, items, total, path, error}
    """
    p = Path(path).expanduser().resolve()
    if not p.exists():
        return {"success": False, "error": f"Diretório não encontrado: {path}"}
    if not p.is_dir():
        return {"success": False, "error": f"Não é um diretório: {path}"}

    items = []
    try:
        if recursive:
            entries = list(p.rglob("*"))
        else:
            entries = list(p.iterdir())

        for entry in sorted(entries)[:max_items]:
            if not show_hidden and entry.name.startswith("."):
                continue
            try:
                stat = entry.stat()
                items.append({
                    "name": entry.name,
                    "path": str(entry),
                    "type": "dir" if entry.is_dir() else "file",
                    "size_bytes": stat.st_size if entry.is_file() else 0,
                    "extension": entry.suffix.lower() if entry.is_file() else "",
                    "modified": stat.st_mtime,
                })
            except (PermissionError, OSError):
                continue

        return {"success": True, "items": items, "total": len(items), "path": str(p)}
    except PermissionError:
        return {"success": False, "error": f"Sem permissão para aceder a: {path}"}


def search_files(base_path: str, pattern: str = "", content_query: str = "",
                 extensions: list = None, max_results: int = 50) -> dict:
    """Search for files by name pattern and/or content.

    Returns: {success, results, total, error}
    """
    base = Path(base_path).expanduser().resolve()
    if not base.exists():
        return {"success": False, "error": f"Caminho não encontrado: {base_path}"}

    results = []
    exts = {e.lower() if e.startswith(".") else f".{e.lower()}" for e in (extensions or [])}

    try:
        glob_pattern = f"**/*{pattern}*" if pattern else "**/*"
        for entry in base.rglob(glob_pattern if pattern else "*"):
            if len(results) >= max_results:
                break
            if not entry.is_file():
                continue
            if exts and entry.suffix.lower() not in exts:
                continue
            if entry.name.startswith("."):
                continue

            match = {"path": str(entry), "name": entry.name, "size_bytes": entry.stat().st_size}

            if content_query:
                try:
                    text = entry.read_text(encoding="utf-8", errors="replace")
                    if content_query.lower() in text.lower():
                        idx = text.lower().find(content_query.lower())
                        match["excerpt"] = text[max(0, idx - 60): idx + 100]
                    else:
                        continue
                except Exception:
                    continue

            results.append(match)

        return {"success": True, "results": results, "total": len(results)}
    except Exception as e:
        return {"success": False, "error": str(e)}


# ── Private readers ────────────────────────────────────────────────────────────

def _read_pdf(path: Path) -> str:
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            pages = []
            for i, page in enumerate(pdf.pages[:20]):
                text = page.extract_text() or ""
                if text:
                    pages.append(f"[Página {i+1}]\n{text}")
            return "\n\n".join(pages)
    except ImportError:
        return f"[pdfplumber não instalado — instala com: pip install pdfplumber]\nCaminho: {path}"


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except ImportError:
        return f"[python-docx não instalado — instala com: pip install python-docx]\nCaminho: {path}"


def _read_excel(path: Path) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        parts = []
        for sheet_name in wb.sheetnames[:5]:
            sheet = wb[sheet_name]
            rows = []
            for row in sheet.iter_rows(max_row=50, values_only=True):
                row_text = "\t".join(str(c) if c is not None else "" for c in row)
                if row_text.strip():
                    rows.append(row_text)
            if rows:
                parts.append(f"[Folha: {sheet_name}]\n" + "\n".join(rows))
        return "\n\n".join(parts)
    except ImportError:
        return f"[openpyxl não instalado]\nCaminho: {path}"


def _read_yaml(path: Path) -> str:
    try:
        import yaml
        with open(path, encoding="utf-8") as f:
            data = yaml.safe_load(f)
        return json.dumps(data, ensure_ascii=False, indent=2)
    except ImportError:
        return path.read_text(encoding="utf-8", errors="replace")
